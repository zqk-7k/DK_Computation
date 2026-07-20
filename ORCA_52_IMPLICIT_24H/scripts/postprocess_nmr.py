#!/usr/bin/env python3
"""Convert ORCA shieldings into auditable H/C NMR prediction products.

Chemical shifts are DFT-derived.  The time-bounded proton and heteronuclear
splitting model is empirical and explicitly labelled.  Synthetic noise,
baseline, water/HOD and solvent peaks are written separately from the clean
calculated spectrum.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
from collections import Counter, defaultdict, deque
from pathlib import Path
from typing import Any, Iterable, Sequence

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from rdkit import Chem

from common import (
    ROOT,
    read_csv,
    read_json,
    read_xyz,
    runtime_environment_identity,
    sha256,
    write_csv,
    write_json,
)


PREDICTION_LABEL = "CALCULATED/SYNTHETIC - NOT EXPERIMENTAL"
SHIFT_LABEL = "DFT CHEMICAL SHIFTS"
J_LABEL = "EMPIRICAL GEOMETRY/TOPOLOGY J - NOT ORCA J"


def truth(value: str) -> bool:
    return value.strip().lower() in {"1", "true", "yes", "y"}


def method_table() -> dict[str, dict[str, Any]]:
    config = read_json(ROOT / "config" / "methods_24h.json")
    return {row["method_id"]: row for row in config["shielding_methods"]}


def validated_parsed_rows(run_dir: Path, path: Path) -> list[dict[str, str]] | None:
    run_config_path = run_dir / "run_config.json"
    if not run_config_path.is_file():
        return None
    identity = read_json(run_config_path)
    expected_uuid = str(identity.get("run_uuid", ""))
    expected_fingerprint = str(identity.get("workflow_fingerprint", ""))
    rows = read_csv(path)
    if not rows or not expected_uuid or not expected_fingerprint:
        return None
    if any(
        row.get("run_uuid") != expected_uuid
        or row.get("workflow_fingerprint") != expected_fingerprint
        or not row.get("source_input_sha256")
        or not row.get("source_geometry_sha256")
        for row in rows
    ):
        return None
    first = rows[0]
    relative = first.get("source_output_relative_to_run_dir", "")
    raw_output = run_dir / relative if relative else Path(first.get("source_output", ""))
    if not raw_output.is_file() or sha256(raw_output) != first.get("source_output_sha256"):
        return None
    if any(row.get("source_output_sha256") != first.get("source_output_sha256") for row in rows):
        return None
    return rows


def select_prediction(run_dir: Path, row: dict[str, str], methods: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
    folder = run_dir / "parsed_shieldings" / "compounds" / row["figure_id"]
    candidates: list[dict[str, Any]] = []
    for path in folder.glob("*.csv") if folder.is_dir() else []:
        method_id = path.stem
        if method_id not in methods:
            continue
        parsed = validated_parsed_rows(run_dir, path)
        if not parsed:
            continue
        candidates.append(
            {
                "method_id": method_id,
                "quality_rank": int(methods[method_id]["quality_rank"]),
                "quality_label": methods[method_id]["quality_label"],
                "geometry_tier": parsed[0]["geometry_tier"],
                "path": path,
                "rows": parsed,
            }
        )
    if not candidates:
        return None
    return max(candidates, key=lambda item: item["quality_rank"])


def reference_id_for(row: dict[str, str]) -> str:
    return "TMS_DMSO" if row["orca_continuum_keyword"] == "DMSO" else "TSP_MINUS_WATER"


def reference_means(run_dir: Path, reference_id: str, method_id: str) -> dict[str, float]:
    shield_path = run_dir / "parsed_shieldings" / "references" / reference_id / f"{method_id}.csv"
    shield = validated_parsed_rows(run_dir, shield_path)
    if not shield:
        raise ValueError(f"Reference shielding provenance is invalid: {shield_path}")
    atom_map = {
        int(row["atom_index_1based"]): row
        for row in read_csv(ROOT / "references" / reference_id / "atom_map.csv")
    }
    values: dict[str, list[float]] = defaultdict(list)
    for row in shield:
        index = int(row["atom_index_1based"])
        if truth(atom_map[index].get("reference_signal", "false")):
            values[row["element"]].append(float(row["isotropic_shielding_ppm"]))
    means = {element: float(np.mean(items)) for element, items in values.items()}
    if not {"H", "C"}.issubset(means):
        raise ValueError(f"Reference {reference_id}/{method_id} lacks H or C zero-shift nuclei")
    return means


def canonical_ranks(smiles: str) -> list[int]:
    mol = Chem.AddHs(Chem.MolFromSmiles(smiles))
    return list(Chem.CanonicalRankAtoms(mol, breakTies=False, includeChirality=True))


def build_signal_groups(
    row: dict[str, str],
    selected: dict[str, Any],
    ref_means: dict[str, float],
) -> tuple[list[dict[str, Any]], dict[int, int]]:
    atom_map_rows = read_csv(ROOT / "structures" / row["figure_id"] / "atom_map.csv")
    atom_map = {int(atom["atom_index_1based"]): atom for atom in atom_map_rows}
    ranks = canonical_ranks(row["model_smiles"])
    shielding = {int(item["atom_index_1based"]): item for item in selected["rows"]}
    grouped: dict[tuple[str, int], list[int]] = defaultdict(list)
    for index, atom in atom_map.items():
        element = atom["element"]
        if element not in {"H", "C"}:
            continue
        if element == "H" and not truth(atom["include_in_analyte_nmr"]):
            continue
        grouped[(element, ranks[index - 1])].append(index)

    signals: list[dict[str, Any]] = []
    atom_to_signal: dict[int, int] = {}
    for signal_index, ((element, rank), indices) in enumerate(
        sorted(grouped.items(), key=lambda item: (item[0][0], min(item[1]))), start=1
    ):
        shields = [float(shielding[index]["isotropic_shielding_ppm"]) for index in indices]
        deltas = [ref_means[element] - value for value in shields]
        signal = {
            "signal_index": signal_index,
            "element": element,
            "canonical_symmetry_rank": rank,
            "atom_indices_1based": indices,
            "atom_indices_text": ";".join(str(index) for index in indices),
            "shielding_mean_ppm": float(np.mean(shields)),
            "delta_ppm": float(np.mean(deltas)),
            "integral_or_equivalent_atom_count": len(indices),
            "exchangeable": any(truth(atom_map[index]["exchangeable"]) for index in indices),
            "method_id": selected["method_id"],
            "method_quality_label": selected["quality_label"],
            "geometry_tier": selected["geometry_tier"],
            "reference_id": reference_id_for(row),
            "reference_shielding_ppm": ref_means[element],
        }
        signals.append(signal)
        for index in indices:
            atom_to_signal[index] = signal_index
    return signals, atom_to_signal


def adjacency_and_bonds(figure_id: str) -> tuple[dict[int, set[int]], dict[tuple[int, int], dict[str, Any]]]:
    adjacency: dict[int, set[int]] = defaultdict(set)
    bonds: dict[tuple[int, int], dict[str, Any]] = {}
    for row in read_csv(ROOT / "structures" / figure_id / "bonds.csv"):
        i, j = int(row["atom_i_1based"]), int(row["atom_j_1based"])
        adjacency[i].add(j)
        adjacency[j].add(i)
        bonds[tuple(sorted((i, j)))] = {
            "order": float(row["bond_order"]),
            "aromatic": truth(row["is_aromatic"]),
        }
    return adjacency, bonds


def shortest_path(adjacency: dict[int, set[int]], start: int, end: int, max_edges: int = 6) -> list[int] | None:
    queue: deque[list[int]] = deque([[start]])
    visited = {start}
    while queue:
        path = queue.popleft()
        if len(path) - 1 >= max_edges:
            continue
        for neighbor in adjacency[path[-1]]:
            if neighbor == end:
                return [*path, neighbor]
            if neighbor not in visited:
                visited.add(neighbor)
                queue.append([*path, neighbor])
    return None


def dihedral_degrees(points: Sequence[np.ndarray]) -> float:
    p0, p1, p2, p3 = points
    b0 = -(p1 - p0)
    b1 = p2 - p1
    b2 = p3 - p2
    b1 = b1 / np.linalg.norm(b1)
    v = b0 - np.dot(b0, b1) * b1
    w = b2 - np.dot(b2, b1) * b1
    x = np.dot(v, w)
    y = np.dot(np.cross(b1, v), w)
    return float(np.degrees(np.arctan2(y, x)))


def empirical_hh_j(
    path: list[int],
    elements: dict[int, str],
    coords: dict[int, np.ndarray],
    bonds: dict[tuple[int, int], dict[str, Any]],
) -> float:
    edges = len(path) - 1
    if edges == 2:
        center = elements[path[1]]
        return 12.0 if center == "C" else 5.5
    if edges == 3:
        middle_bond = bonds.get(tuple(sorted((path[1], path[2]))), {})
        if middle_bond.get("aromatic"):
            return 8.0
        if elements[path[1]] == "C" and elements[path[2]] == "C":
            theta = math.radians(dihedral_degrees([coords[index] for index in path]))
            return max(0.5, min(14.0, 12.0 * math.cos(theta) ** 2 - math.cos(theta) + 0.5))
        return 6.5
    if edges == 4 and all(
        bonds.get(tuple(sorted((a, b))), {}).get("aromatic") for a, b in zip(path[1:-2], path[2:-1])
    ):
        return 2.0
    if edges == 5:
        return 0.6
    return 0.0


def empirical_hx_j(path: list[int], hetero: str) -> float:
    edges = len(path) - 1
    if hetero == "F":
        return {1: 50.0, 2: 45.0, 3: 9.0, 4: 3.0}.get(edges, 0.0)
    if hetero == "P":
        return {1: 180.0, 2: 15.0, 3: 7.0, 4: 2.0}.get(edges, 0.0)
    return 0.0


def empirical_cx_j(path: list[int], hetero: str) -> float:
    edges = len(path) - 1
    if hetero == "F":
        return {1: 250.0, 2: 22.0, 3: 8.0, 4: 3.0}.get(edges, 0.0)
    if hetero == "P":
        return {1: 100.0, 2: 18.0, 3: 7.0, 4: 2.0}.get(edges, 0.0)
    return 0.0


def coupling_models(
    row: dict[str, str],
    signals: list[dict[str, Any]],
    atom_to_signal: dict[int, int],
    geometry: Path,
) -> dict[int, list[dict[str, Any]]]:
    atom_map_rows = read_csv(ROOT / "structures" / row["figure_id"] / "atom_map.csv")
    elements = {int(atom["atom_index_1based"]): atom["element"] for atom in atom_map_rows}
    exchangeable = {int(atom["atom_index_1based"]): truth(atom["exchangeable"]) for atom in atom_map_rows}
    _, atoms = read_xyz(geometry)
    coords = {index: np.array(atom[1:], dtype=float) for index, atom in enumerate(atoms, start=1)}
    adjacency, bonds = adjacency_and_bonds(row["figure_id"])
    result: dict[int, list[dict[str, Any]]] = defaultdict(list)

    h_signals = [signal for signal in signals if signal["element"] == "H"]
    for signal in h_signals:
        sid = int(signal["signal_index"])
        if signal["exchangeable"]:
            continue
        representative = int(signal["atom_indices_1based"][0])
        partner_groups: dict[tuple[str, int, float], set[int]] = defaultdict(set)
        for other in h_signals:
            oid = int(other["signal_index"])
            if oid == sid or other["exchangeable"]:
                continue
            for partner_atom in other["atom_indices_1based"]:
                path = shortest_path(adjacency, representative, int(partner_atom), max_edges=5)
                if not path:
                    continue
                j = empirical_hh_j(path, elements, coords, bonds)
                if abs(j) >= 0.4:
                    partner_groups[("H", oid, round(abs(j), 1))].add(int(partner_atom))
        for atom_index, element in elements.items():
            if element not in {"F", "P"}:
                continue
            path = shortest_path(adjacency, representative, atom_index, max_edges=4)
            if not path:
                continue
            j = empirical_hx_j(path, element)
            if abs(j) >= 0.4:
                partner_groups[(element, atom_index, round(abs(j), 1))].add(atom_index)
        entries = []
        for (partner_element, partner_id, j), atoms_set in partner_groups.items():
            # Count only partner nuclei that have this particular topology/J.
            # A symmetry group can contain both ortho- and meta-coupled nuclei;
            # using the full group integral for each J would over-split A2B2 rings.
            count = len(atoms_set)
            entries.append(
                {
                    "partner_type": partner_element,
                    "partner_signal_or_atom": partner_id,
                    "equivalent_partner_count": count,
                    "J_Hz": j,
                    "model": J_LABEL,
                }
            )
        result[sid] = sorted(entries, key=lambda item: abs(item["J_Hz"]), reverse=True)[:6]

    hetero_atoms = [(index, element) for index, element in elements.items() if element in {"F", "P"}]
    for signal in [item for item in signals if item["element"] == "C"]:
        sid = int(signal["signal_index"])
        representative = int(signal["atom_indices_1based"][0])
        groups: dict[tuple[str, float], int] = defaultdict(int)
        for atom_index, element in hetero_atoms:
            path = shortest_path(adjacency, representative, atom_index, max_edges=4)
            if not path:
                continue
            j = empirical_cx_j(path, element)
            if abs(j) >= 0.4:
                groups[(element, round(abs(j), 1))] += 1
        result[sid] = [
            {
                "partner_type": element,
                "partner_signal_or_atom": "heteronuclear",
                "equivalent_partner_count": count,
                "J_Hz": j,
                "model": J_LABEL,
            }
            for (element, j), count in sorted(groups.items(), key=lambda item: item[0][1], reverse=True)
        ][:5]
    return result


def pascal(n: int) -> list[int]:
    return [math.comb(n, k) for k in range(n + 1)]


def first_order_lines(
    center_ppm: float,
    intensity: float,
    couplings: list[dict[str, Any]],
    spectrometer_mhz: float,
    max_lines: int = 512,
) -> list[tuple[float, float]]:
    def merge_nearby(items: list[tuple[float, float]], tolerance_ppm: float) -> list[tuple[float, float]]:
        merged: list[tuple[float, float]] = []
        for ppm, area in sorted(items):
            if merged and abs(ppm - merged[-1][0]) <= tolerance_ppm:
                old_ppm, old_area = merged[-1]
                total = old_area + area
                merged[-1] = ((old_ppm * old_area + ppm * area) / total, total)
            else:
                merged.append((ppm, area))
        return merged

    def compress_moment_preserving(items: list[tuple[float, float]], limit: int) -> list[tuple[float, float]]:
        if len(items) <= limit:
            return items
        ordered = sorted(items)
        compressed: list[tuple[float, float]] = []
        for chunk in np.array_split(np.asarray(ordered, dtype=float), limit):
            areas = chunk[:, 1]
            total = float(np.sum(areas))
            if total <= 0:
                continue
            compressed.append((float(np.sum(chunk[:, 0] * areas) / total), total))
        return compressed

    lines = [(center_ppm, float(intensity))]
    tolerance_ppm = 0.05 / spectrometer_mhz
    for coupling in couplings:
        count = max(1, min(3, int(coupling["equivalent_partner_count"])))
        j_ppm = float(coupling["J_Hz"]) / spectrometer_mhz
        weights = pascal(count)
        new_lines: list[tuple[float, float]] = []
        for ppm, amplitude in lines:
            for k, weight in enumerate(weights):
                offset = (k - count / 2.0) * j_ppm
                new_lines.append((ppm + offset, amplitude * weight / sum(weights)))
        lines = merge_nearby(new_lines, tolerance_ppm)
        lines = compress_moment_preserving(lines, max_lines)
    lines = merge_nearby(lines, tolerance_ppm)
    lines = compress_moment_preserving(lines, max_lines)
    total_area = sum(area for _, area in lines)
    if total_area <= 0:
        return [(center_ppm, float(intensity))]
    weighted_center = sum(ppm * area for ppm, area in lines) / total_area
    center_correction = center_ppm - weighted_center
    scale = float(intensity) / total_area
    return [(ppm + center_correction, area * scale) for ppm, area in sorted(lines)]


def lorentzian_area(x: np.ndarray, center: float, fwhm_ppm: float, area: float) -> np.ndarray:
    half = max(fwhm_ppm / 2.0, 1.0e-7)
    return (area / math.pi) * half / ((x - center) ** 2 + half * half)


def lorentzian_height(x: np.ndarray, center: float, fwhm_ppm: float, height: float) -> np.ndarray:
    half = max(fwhm_ppm / 2.0, 1.0e-7)
    return height * half * half / ((x - center) ** 2 + half * half)


def synthesize_profile(
    signals: list[dict[str, Any]],
    couplings: dict[int, list[dict[str, Any]]],
    nucleus: str,
    solvent_label: str,
    config: dict[str, Any],
    seed: int,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, list[dict[str, Any]], dict[str, Any]]:
    preset = config[nucleus]
    frequency = float(preset["spectrometer_MHz"])
    points = int(preset["points"])
    x = np.linspace(float(preset["ppm_high"]), float(preset["ppm_low"]), points)
    clean = np.zeros_like(x)
    appearance = np.zeros_like(x)
    rng = np.random.default_rng(seed + (1 if nucleus == "1H" else 13))
    line_rows: list[dict[str, Any]] = []

    for signal in [item for item in signals if item["element"] == ("H" if nucleus == "1H" else "C")]:
        sid = int(signal["signal_index"])
        center = float(signal["delta_ppm"])
        integral = float(signal["integral_or_equivalent_atom_count"])
        signal_couplings = couplings.get(sid, [])
        lines = first_order_lines(center, integral, signal_couplings, frequency)
        if nucleus == "1H":
            clean_fwhm = 12.0 if signal["exchangeable"] else 1.0
            median = float(
                preset["exchangeable_fwhm_Hz_median"]
                if signal["exchangeable"] else preset["nonexchangeable_fwhm_Hz_median"]
            )
            sigma = float(
                preset["exchangeable_fwhm_log_sigma"]
                if signal["exchangeable"] else preset["nonexchangeable_fwhm_log_sigma"]
            )
        else:
            clean_fwhm = 1.0
            median = float(preset["analyte_fwhm_Hz_median"])
            sigma = float(preset["analyte_fwhm_log_sigma"])
        synthetic_fwhm = float(rng.lognormal(math.log(median), sigma))
        signal["clean_fwhm_Hz"] = clean_fwhm
        signal["synthetic_fwhm_Hz"] = synthetic_fwhm
        signal["couplings_json"] = json.dumps(signal_couplings, ensure_ascii=False, separators=(",", ":"))
        for line_index, (ppm, amplitude) in enumerate(lines, start=1):
            clean += lorentzian_area(x, ppm, clean_fwhm / frequency, amplitude)
            appearance += lorentzian_area(x, ppm, synthetic_fwhm / frequency, amplitude)
            line_rows.append(
                {
                    "signal_index": sid,
                    "line_index": line_index,
                    "ppm": ppm,
                    "relative_area": amplitude,
                    "clean_fwhm_Hz": clean_fwhm,
                    "synthetic_fwhm_Hz": synthetic_fwhm,
                }
            )

    analyte_max = float(np.max(appearance)) if appearance.size and np.max(appearance) > 0 else 1.0
    inserted: list[dict[str, Any]] = []
    if nucleus == "1H":
        if solvent_label == "DMSO-d6":
            solvent = preset["DMSO-d6"]
            solvent_height = float(rng.uniform(*solvent["solvent_height_range"])) * analyte_max
            water_ratio = float(np.exp(rng.uniform(math.log(solvent["water_to_solvent_height_range"][0]), math.log(solvent["water_to_solvent_height_range"][1]))))
            water_height = water_ratio * solvent_height
            appearance += lorentzian_height(
                x, float(solvent["residual_solvent_ppm"]),
                float(solvent["residual_solvent_fwhm_Hz"]) / frequency, solvent_height,
            )
            appearance += lorentzian_height(
                x, float(solvent["water_ppm"]),
                float(solvent["water_fwhm_Hz"]) / frequency, water_height,
            )
            inserted.extend(
                [
                    {"component": "residual_DMSO", "ppm": solvent["residual_solvent_ppm"], "height": solvent_height},
                    {"component": "water_in_DMSO", "ppm": solvent["water_ppm"], "height": water_height},
                ]
            )
        else:
            solvent = preset["D2O"]
            hod_height = float(rng.uniform(*solvent["HOD_height_range"])) * analyte_max
            hod_width = float(rng.uniform(*solvent["HOD_fwhm_Hz_range"]))
            appearance += lorentzian_height(x, float(solvent["HOD_ppm"]), hod_width / frequency, hod_height)
            inserted.append({"component": "HOD", "ppm": solvent["HOD_ppm"], "height": hod_height})
    elif solvent_label == "DMSO-d6":
        solvent = preset["DMSO-d6"]
        central = float(solvent["solvent_central_ppm"])
        spacing_ppm = float(solvent["septet_spacing_Hz"]) / frequency
        solvent_height = float(rng.uniform(*solvent["solvent_height_range"])) * analyte_max
        weights = np.array([1, 3, 6, 7, 6, 3, 1], dtype=float) / 7.0
        width = float(rng.uniform(*solvent["component_fwhm_Hz_range"]))
        for index, weight in enumerate(weights):
            ppm = central + (index - 3) * spacing_ppm
            appearance += lorentzian_height(x, ppm, width / frequency, solvent_height * float(weight))
        inserted.append({"component": "DMSO-d6_13C_septet", "ppm": central, "height": solvent_height})

    scale = float(np.max(np.abs(appearance))) if np.max(np.abs(appearance)) > 0 else 1.0
    noise_fraction = float(rng.uniform(*preset["noise_sigma_fraction_range"]))
    white = rng.normal(0.0, noise_fraction * scale, points)
    correlated = np.convolve(white, np.array([0.18, 0.64, 0.18]), mode="same")
    baseline_fraction = float(rng.uniform(*preset["baseline_fraction_range"]))
    u = np.linspace(-1.0, 1.0, points)
    coefficients = rng.normal(0.0, baseline_fraction * scale, 4)
    baseline = coefficients[0] + coefficients[1] * u + coefficients[2] * u**2 + coefficients[3] * u**3
    baseline += baseline_fraction * scale * 0.2 * np.sin(2 * np.pi * (u + 1) * rng.uniform(0.5, 2.0))
    appearance = appearance + correlated + baseline
    metadata = {
        "nucleus": nucleus,
        "spectrometer_MHz": frequency,
        "point_count": points,
        "noise_sigma_fraction": noise_fraction,
        "baseline_fraction": baseline_fraction,
        "inserted_background_components": inserted,
        "random_seed": seed + (1 if nucleus == "1H" else 13),
        "appearance_label": "SYNTHETIC_APPEARANCE_NOT_ACQUIRED",
    }
    target_area = float(sum(row["relative_area"] for row in line_rows))
    numeric_area = float(abs(np.trapezoid(clean, x)))
    metadata["analyte_target_area"] = target_area
    metadata["clean_numeric_area_within_display_window"] = numeric_area
    metadata["clean_area_capture_fraction"] = numeric_area / target_area if target_area > 0 else 1.0
    if target_area > 0 and numeric_area > 0:
        # The NMR axis is descending, so both signed trapezoidal integrals are
        # negative.  Taking the ratio preserves the physical centroid without
        # accidentally returning a negative ppm value.
        denominator = float(np.trapezoid(clean, x))
        numerator = float(np.trapezoid(clean * x, x))
        metadata["clean_centroid_ppm"] = numerator / denominator
    return x, clean, appearance, line_rows, metadata


def save_xy(folder: Path, nucleus: str, x: np.ndarray, clean: np.ndarray, appearance: np.ndarray) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    rows = [
        {"ppm": float(ppm), "clean_calculated_intensity": float(c), "synthetic_appearance_intensity": float(a)}
        for ppm, c, a in zip(x, clean, appearance)
    ]
    write_csv(folder / f"{nucleus}_spectrum_full_points.csv", rows, list(rows[0].keys()))
    with (folder / f"{nucleus}_spectrum_full_points.txt").open("w", encoding="utf-8", newline="") as handle:
        handle.write(f"# {PREDICTION_LABEL}\n")
        handle.write("# ppm\tclean_calculated_intensity\tsynthetic_appearance_intensity\n")
        for row in rows:
            handle.write(
                f"{row['ppm']:.9f}\t{row['clean_calculated_intensity']:.10e}\t"
                f"{row['synthetic_appearance_intensity']:.10e}\n"
            )


def plot_spectrum(
    path: Path,
    figure_id: str,
    display_name: str,
    nucleus: str,
    x: np.ndarray,
    y: np.ndarray,
    appearance: bool,
    status_note: str,
) -> None:
    plt.rcParams.update({"font.family": "Arial", "font.size": 8, "axes.linewidth": 0.8})
    fig, ax = plt.subplots(figsize=(8.0, 3.25))
    ax.plot(x, y, color="#171717", linewidth=0.72)
    ax.set_xlim(float(x[0]), float(x[-1]))
    ax.set_xlabel(f"{nucleus} chemical shift (ppm)")
    ax.set_yticks([])
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_visible(False)
    mode = "synthetic acquisition appearance" if appearance else "clean calculated profile"
    ax.set_title(f"{figure_id} | {display_name} | {nucleus} {mode}", loc="left", fontsize=9, fontweight="bold")
    fig.text(0.99, 0.965, PREDICTION_LABEL, ha="right", va="top", fontsize=8, color="#9B1C1C", fontweight="bold")
    fig.text(0.01, 0.02, status_note, ha="left", va="bottom", fontsize=6.5, color="#555555")
    fig.subplots_adjust(left=0.055, right=0.99, top=0.80, bottom=0.20)
    fig.savefig(path.with_suffix(".png"), dpi=300, metadata={"Title": PREDICTION_LABEL})
    fig.savefig(path.with_suffix(".svg"), metadata={"Title": PREDICTION_LABEL})
    plt.close(fig)


def write_peak_tables(folder: Path, nucleus: str, signals: list[dict[str, Any]], line_rows: list[dict[str, Any]]) -> None:
    selected = [signal for signal in signals if signal["element"] == ("H" if nucleus == "1H" else "C")]
    fields = [
        "signal_index", "element", "canonical_symmetry_rank", "atom_indices_text",
        "shielding_mean_ppm", "delta_ppm", "integral_or_equivalent_atom_count",
        "exchangeable", "clean_fwhm_Hz", "synthetic_fwhm_Hz", "couplings_json",
        "method_id", "method_quality_label", "geometry_tier", "reference_id",
        "reference_shielding_ppm",
    ]
    normalized = []
    for signal in selected:
        normalized.append({key: signal.get(key, "") for key in fields})
    if normalized:
        write_csv(folder / f"{nucleus}_predicted_peaks.csv", normalized, fields)
    else:
        write_csv(
            folder / f"{nucleus}_predicted_peaks.csv",
            [{"status": "NO_PERSISTENT_ANALYTE_SIGNAL_EXPECTED"}],
            ["status"],
        )
    if line_rows:
        write_csv(folder / f"{nucleus}_first_order_lines.csv", line_rows, list(line_rows[0].keys()))
    else:
        write_csv(folder / f"{nucleus}_first_order_lines.csv", [{"status": "NO_ANALYTE_LINES"}], ["status"])


def write_workbook(folder: Path, metadata: dict[str, Any]) -> None:
    workbook = Workbook(write_only=True)
    meta = workbook.create_sheet("README")
    meta.append([PREDICTION_LABEL])
    meta.append(["key", "value"])
    for key, value in metadata.items():
        meta.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    for nucleus in ("1H", "13C"):
        peak_path = folder / f"{nucleus}_predicted_peaks.csv"
        sheet = workbook.create_sheet(f"{nucleus}_peaks")
        with peak_path.open("r", encoding="utf-8-sig", newline="") as handle:
            for row in csv.reader(handle):
                sheet.append(row)
        spectrum_path = folder / f"{nucleus}_spectrum_full_points.csv"
        sheet = workbook.create_sheet(f"{nucleus}_spectrum")
        with spectrum_path.open("r", encoding="utf-8-sig", newline="") as handle:
            for row in csv.reader(handle):
                sheet.append(row)
    workbook.save(folder / "NMR_prediction_data.xlsx")


def geometry_for_selected(run_dir: Path, figure_id: str, selected: dict[str, Any]) -> Path:
    if selected["geometry_tier"] == "REFINED_GEOMETRY":
        candidate = run_dir / "geometries" / "refined" / f"{figure_id}.xyz"
        if candidate.is_file():
            return candidate
    return run_dir / "geometries" / "baseline" / f"{figure_id}.xyz"


def empty_signals_profile(
    row: dict[str, str], config: dict[str, Any], seed: int, nucleus: str
) -> tuple[np.ndarray, np.ndarray, np.ndarray, list[dict[str, Any]], dict[str, Any]]:
    return synthesize_profile([], {}, nucleus, row["experimental_solvent_label"], config, seed)


def process_compound(
    run_dir: Path,
    row: dict[str, str],
    methods: dict[str, dict[str, Any]],
    appearance_config: dict[str, Any],
) -> dict[str, Any]:
    figure_id = row["figure_id"]
    folder = run_dir / "NMR_CALCULATED_NOT_EXPERIMENTAL" / figure_id
    folder.mkdir(parents=True, exist_ok=True)
    seed = int(hashlib.sha256(figure_id.encode()).hexdigest()[:8], 16)
    selected = select_prediction(run_dir, row, methods)
    no_h = int(row["observable_1H_count_after_exchange_policy"]) == 0
    no_c = int(row["observable_13C_count"]) == 0
    if selected is None and not (no_h and no_c):
        result = {
            "figure_id": figure_id,
            "display_name": row["display_name"],
            "status": "MISSING_ORCA_SHIELDING",
            "observable_1H": not no_h,
            "observable_13C": not no_c,
            "assumed_pd": row["assumed_pd"],
            "species_note": row["species_note"],
            "selected_source_scenario": row["selected_source_scenario"],
        }
        write_json(folder / "metadata.json", result)
        return result

    if selected is not None:
        ref_id = reference_id_for(row)
        refs = reference_means(run_dir, ref_id, selected["method_id"])
        reference_shielding_path = (
            run_dir / "parsed_shieldings" / "references" / ref_id / f"{selected['method_id']}.csv"
        )
        signals, atom_to_signal = build_signal_groups(row, selected, refs)
        geometry = geometry_for_selected(run_dir, figure_id, selected)
        couplings = coupling_models(row, signals, atom_to_signal, geometry)
    else:
        refs = {}
        reference_shielding_path = None
        signals = []
        couplings = {}
        geometry = ROOT / "structures" / figure_id / "start.xyz"

    nucleus_products: dict[str, Any] = {}
    for nucleus in ("1H", "13C"):
        if nucleus == "13C" and no_c:
            preset = appearance_config["13C"]
            x = np.linspace(float(preset["ppm_high"]), float(preset["ppm_low"]), int(preset["points"]))
            clean = np.zeros_like(x)
            appearance = np.zeros_like(x)
            line_rows = []
            profile_meta = {
                "nucleus": "13C",
                "status": "NOT_APPLICABLE_NO_CARBON",
                "point_count": int(preset["points"]),
                "appearance_label": "NO_SYNTHETIC_TRACE_GENERATED",
            }
        else:
            x, clean, appearance, line_rows, profile_meta = synthesize_profile(
                signals, couplings, nucleus, row["experimental_solvent_label"], appearance_config, seed
            )
        save_xy(folder, nucleus, x, clean, appearance)
        write_peak_tables(folder, nucleus, signals, line_rows)
        if nucleus == "1H" and no_h:
            status_note = "NO_PERSISTENT_ANALYTE_1H_EXPECTED; D2O exchange policy; appearance contains HOD/background only"
        elif nucleus == "13C" and no_c:
            status_note = "NOT_APPLICABLE_NO_CARBON; no analyte carbon peak is generated"
        else:
            status_note = f"{SHIFT_LABEL}; {J_LABEL}; {selected['method_id'] if selected else 'N/A'}"
        plot_spectrum(folder / f"{nucleus}_clean", figure_id, row["display_name"], nucleus, x, clean, False, status_note)
        plot_spectrum(folder / f"{nucleus}_synthetic_appearance", figure_id, row["display_name"], nucleus, x, appearance, True, status_note)
        nucleus_products[nucleus] = profile_meta

    outlier_flags = []
    for signal in signals:
        delta = float(signal["delta_ppm"])
        if signal["element"] == "H" and not (-2.0 <= delta <= 18.0):
            outlier_flags.append(f"H signal {signal['signal_index']} delta={delta:.3f} ppm")
        if signal["element"] == "C" and not (-30.0 <= delta <= 260.0):
            outlier_flags.append(f"C signal {signal['signal_index']} delta={delta:.3f} ppm")
    if selected is None:
        conformer_treatment = "NOT_APPLICABLE"
    elif (run_dir / "geometries" / "goat" / f"{figure_id}.xyz").is_file():
        conformer_treatment = "SINGLE_CONFORMER_GOAT_GLOBAL_MINIMUM_APPROXIMATION_NO_BOLTZMANN_ENSEMBLE"
    elif selected["geometry_tier"] == "REFINED_GEOMETRY":
        conformer_treatment = "SINGLE_CONFORMER_R2SCAN_REFINED_APPROXIMATION_NO_BOLTZMANN_ENSEMBLE"
    else:
        conformer_treatment = "SINGLE_CONFORMER_BASELINE_APPROXIMATION_NO_BOLTZMANN_ENSEMBLE"
    salt_model_notice = ""
    if figure_id in {"F3_016", "F3_017"}:
        salt_model_notice = (
            "CATION-ONLY CONTINUUM SENSITIVITY MODEL; CHLORIDE AND ION-PAIR EFFECTS ARE NOT MODELLED"
        )
    elif figure_id == "F3_021":
        salt_model_notice = (
            "DICATION-ONLY CONTINUUM MODEL; TWO CHLORIDE COUNTERIONS AND CONTACT-ION-PAIR EFFECTS ARE NOT MODELLED"
        )
    metadata = {
        "schema_version": "1.0",
        "prediction_label": PREDICTION_LABEL,
        "figure_id": figure_id,
        "display_name": row["display_name"],
        "model_smiles": row["model_smiles"],
        "model_formula": row["model_formula"],
        "charge": int(row["charge"]),
        "multiplicity": int(row["multiplicity"]),
        "experimental_solvent_label": row["experimental_solvent_label"],
        "orca_continuum_keyword": row["orca_continuum_keyword"],
        "assumed_pd": row["assumed_pd"],
        "species_note": row["species_note"],
        "selected_source_scenario": row["selected_source_scenario"],
        "species_model_status": row["status"],
        "salt_continuum_notice": salt_model_notice,
        "heavy_atom_method_notice": (
            "IODINE_USES_DEF2_TZVP_DEF2_ECP; NO_ALL_ELECTRON_X2C_SENSITIVITY_IN_24H_PROFILE; NEAR_I_HC_SHIFTS_HAVE_HIGHER_UNCERTAINTY"
            if figure_id == "F3_006" else "NOT_APPLICABLE"
        ),
        "selected_method_id": selected["method_id"] if selected else "NOT_APPLICABLE",
        "selected_method_quality_label": selected["quality_label"] if selected else "NOT_APPLICABLE",
        "selected_geometry_tier": selected["geometry_tier"] if selected else "NOT_APPLICABLE",
        "conformer_treatment": conformer_treatment,
        "selected_shielding_file": str(selected["path"]) if selected else "NOT_APPLICABLE",
        "selected_shielding_sha256": sha256(selected["path"]) if selected else "NOT_APPLICABLE",
        "reference_shieldings_ppm": refs,
        "reference_shielding_file": str(reference_shielding_path) if reference_shielding_path else "NOT_APPLICABLE",
        "reference_shielding_sha256": sha256(reference_shielding_path) if reference_shielding_path else "NOT_APPLICABLE",
        "geometry": str(geometry),
        "geometry_sha256": sha256(geometry),
        "proton_splitting_model": J_LABEL,
        "carbon_mode": "13C{1H}; empirical F/P splitting retained when topology indicates it",
        "no_persistent_analyte_1H_expected": no_h,
        "no_carbon_13C_not_applicable": no_c,
        "outlier_flags": outlier_flags,
        "profiles": nucleus_products,
        "status": "COMPLETE_WITH_QC_FLAGS" if outlier_flags else "COMPLETE",
    }
    write_json(folder / "metadata.json", metadata)
    write_workbook(folder, metadata)
    return metadata


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--only", default="")
    args = parser.parse_args()
    run_config = read_json(args.run_dir / "run_config.json")
    expected_environment = run_config.get("runtime_environment_identity")
    if expected_environment and expected_environment != runtime_environment_identity():
        raise RuntimeError(
            "The Python/scientific-package environment changed after the ORCA run. "
            "Use the recorded environment or start a new run directory."
        )
    species = read_csv(ROOT / "config" / "primary_species.csv")
    selected_ids = {item.strip() for item in args.only.split(",") if item.strip()}
    species = [row for row in species if not selected_ids or row["figure_id"] in selected_ids]
    methods = method_table()
    appearance = read_json(ROOT / "config" / "appearance_400MHz.json")
    results: list[dict[str, Any]] = []
    for row in species:
        print(f"Post-processing {row['figure_id']} {row['display_name']}", flush=True)
        try:
            results.append(process_compound(args.run_dir, row, methods, appearance))
        except Exception as exc:
            failure = {
                "figure_id": row["figure_id"], "display_name": row["display_name"],
                "status": "POSTPROCESS_EXCEPTION", "error": f"{type(exc).__name__}: {exc}",
            }
            results.append(failure)
            folder = args.run_dir / "NMR_CALCULATED_NOT_EXPERIMENTAL" / row["figure_id"]
            folder.mkdir(parents=True, exist_ok=True)
            write_json(folder / "metadata.json", failure)
    fields = sorted({key for result in results for key in result if not isinstance(result[key], (dict, list))})
    write_csv(
        args.run_dir / "NMR_CALCULATED_NOT_EXPERIMENTAL" / "prediction_summary.csv",
        [{key: result.get(key, "") for key in fields} for result in results], fields,
    )
    failures = [result["figure_id"] for result in results if not str(result.get("status", "")).startswith("COMPLETE")]
    summary = {
        "prediction_label": PREDICTION_LABEL,
        "compound_count": len(results),
        "complete_count": len(results) - len(failures),
        "failures": failures,
        "pass": not failures,
    }
    write_json(args.run_dir / "NMR_CALCULATED_NOT_EXPERIMENTAL" / "postprocess_summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if not failures else 2


if __name__ == "__main__":
    raise SystemExit(main())
